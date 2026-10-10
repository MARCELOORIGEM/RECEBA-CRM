"""Leitura da planilha do Funil, sem servidor.

O formato é o da planilha da operação (LEAD ID, NOME LEAD, ENDEREÇO, BAIRRO,
BD ID, NOME BD, LIDER, DATA VISITA, STATUS, OBS). Os casos aqui são os que
aparecem na vida real: o modelo voltando preenchido, o CSV do Excel em
português (ponto e vírgula, Windows-1252), status escrito em maiúsculas,
datas de vários jeitos e planilhas da versão anterior.
"""
import io
import os
import re
from datetime import date
from pathlib import Path

import pytest

os.environ.setdefault("JWT_SECRET", "teste")

from app import funil  # noqa: E402
from app.planilha_leads import (  # noqa: E402
    COLUNAS, PlanilhaInvalida, _data, _dinheiro, gerar_modelo, interpretar, ler_arquivo,
)

CABECALHO_OPERACAO = ("LEAD ID;NOME LEAD;ENDEREÇO;BAIRRO;BD ID;NOME BD;LIDER;DATA VISITA;"
                      "STATUS;OBS")


def _ler(dados, nome="x.csv"):
    return interpretar(ler_arquivo(dados, nome))


def _csv(*linhas):
    return "\r\n".join(linhas).encode("cp1252")


def test_modelo_tem_as_colunas_da_operacao_na_ordem():
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(gerar_modelo()))["Leads"]
    titulos = [c.value for c in ws[1]]
    assert titulos[:10] == CABECALHO_OPERACAO.split(";")
    assert titulos == [c.titulo for c in COLUNAS]
    assert len(ws.data_validations.dataValidation) == 2  # STATUS e ORIGEM


def test_modelo_volta_valido():
    res = _ler(gerar_modelo(), "modelo.xlsx")
    assert res.erros == []
    a, b = res.validos
    assert (a["codigo_externo"], a["name"], a["bairro"], a["bd_nome"], a["lider"], a["stage"]) == (
        "12345678", "Restaurante 123", "UNIÃO", "JUNIAO", "TARCÍSIO", "nao_localizado")
    assert (b["stage"], b["data_visita"]) == ("reuniao", date(2026, 10, 5))


def test_planilha_da_operacao_em_csv():
    """A linha do print: maiúsculas, acentos, status escrito à mão."""
    res = _ler(_csv(
        CABECALHO_OPERACAO,
        "12345678;Restaurante 123;RUA H, 357 - UNIÃO;UNIÃO;;JUNIAO;TARCÍSIO;;NÃO LOCALIZADO;",
        "12345679;Pizzaria;AV. BRASIL;CENTRO;BD02;MARIANA;TARCÍSIO;05/10/2026;SEGUNDA VISITA;volta 15h",
        "12345680;Lanchonete;RUA X;CENTRO;;;;5/10/26;COLHENDO DADOS;",
    ))
    assert res.erros == []
    assert [v["stage"] for v in res.validos] == ["nao_localizado", "segunda_visita", "colhendo_dados"]
    assert res.validos[0]["endereco"] == "RUA H, 357 - UNIÃO"
    assert res.validos[1]["data_visita"] == res.validos[2]["data_visita"] == date(2026, 10, 5)


@pytest.mark.parametrize("escrito,chave", [
    ("REUNIÃO", "reuniao"), ("reuniao", "reuniao"), ("Reunião agendada", "reuniao"),
    ("Não localizado", "nao_localizado"), ("JÁ É PARCEIRO", "ja_parceiro"),
    ("2a visita", "segunda_visita"), ("ATIVADO", "ativado"),
    # Etapas da versão anterior continuam importando.
    ("Novo", "a_visitar"), ("Em negociação", "reuniao"), ("Ganho", "ativado"),
])
def test_status_escrito_de_varios_jeitos(escrito, chave):
    res = _ler(_csv("NOME LEAD;STATUS", f"X;{escrito}"))
    assert res.validos[0]["stage"] == chave


def test_sem_interesse_exige_motivo_e_fechado_e_ambiguo():
    res = _ler(_csv(
        "NOME LEAD;STATUS;MOTIVO",
        "A;SEM INTERESSE;",
        "B;SEM INTERESSE;Comissão alta",
        "C;Fechado;",
    ))
    assert [v["name"] for v in res.validos] == ["B"]
    motivos = {e["linha"]: e["motivo"] for e in res.erros}
    assert "MOTIVO" in motivos[2]
    # "Fechado" pode ser negócio fechado ou porta fechada: não adivinha.
    assert '"Fechado" não existe' in motivos[4]


@pytest.mark.parametrize("bruto,esperado", [
    ("05/10/2026", date(2026, 10, 5)), ("5/10/26", date(2026, 10, 5)),
    ("2026-10-05", date(2026, 10, 5)), ("", None), (None, None),
])
def test_data_visita(bruto, esperado):
    assert _data(bruto) == esperado


def test_data_ilegivel_vira_erro_na_linha():
    res = _ler(_csv("NOME LEAD;DATA VISITA", "A;amanhã"))
    assert "não é uma data" in res.erros[0]["motivo"]


def test_planilha_antiga_ainda_importa():
    """Cabeçalho da versão anterior do modelo (Nome do estabelecimento, Etapa...)."""
    res = _ler(_csv("Nome do estabelecimento;Origem;Etapa;Valor estimado (R$)",
                    "Pastelaria São João;Prospecção ativa;Em negociação;R$ 3.200,00"))
    v = res.validos[0]
    assert (v["name"], v["source"], v["stage"], v["estimated_value"]) == (
        "Pastelaria São João", "prospeccao", "reuniao", 3200.0)


def test_preenchidos_marca_so_as_celulas_com_valor():
    res = _ler(_csv("LEAD ID;NOME LEAD;BAIRRO;STATUS", "1;A;;REUNIÃO"))
    assert set(res.validos[0]["_preenchidos"]) == {"codigo_externo", "name", "stage"}


def test_funil_atual_exportado_volta_igual():
    leads = [{"codigo_externo": "999", "name": "Lead Exportado", "bairro": "CENTRO",
              "stage": "aguardando_documentos", "data_visita": date(2026, 10, 1),
              "estimated_value": 1500.5}]
    v = _ler(gerar_modelo(leads), "funil.xlsx").validos[0]
    assert (v["codigo_externo"], v["stage"], v["data_visita"], v["estimated_value"]) == (
        "999", "aguardando_documentos", date(2026, 10, 1), 1500.5)


@pytest.mark.parametrize("bruto,esperado", [
    ("3200", 3200.0), ("3.200", 3200.0), ("3.200,50", 3200.5), ("R$ 1.250.000,00", 1250000.0),
    ("3,2 mil", 3200.0), (1500, 1500.0), ("", 0.0), (None, 0.0), ("12.5", 12.5),
])
def test_dinheiro(bruto, esperado):
    assert _dinheiro(bruto) == esperado


def test_linhas_em_branco_e_titulo_em_cima_sao_ignorados():
    res = _ler(b"\n\nNOME LEAD;BAIRRO\nA;X\n;\nB;Y\n")
    assert [v["name"] for v in res.validos] == ["A", "B"] and res.total == 2


def test_colunas_desconhecidas_sao_avisadas_e_nao_quebram():
    res = _ler(b"NOME LEAD;CNPJ;Faturamento anual\nA;123;9\n")
    assert res.validos and res.colunas_ignoradas == ["CNPJ", "Faturamento anual"]


@pytest.mark.parametrize("dados,nome,trecho", [
    (b"", "x.csv", "vazio"),
    (b"Cidade;Telefone\nX;1\n", "x.csv", "NOME LEAD"),
    (b"\xd0\xcf\x11\xe0qualquer", "antigo.xls", ".xls"),
    (b"PK\x03\x04lixo", "quebrado.xlsx", "Excel"),
])
def test_arquivo_invalido_explica_o_que_fazer(dados, nome, trecho):
    with pytest.raises(PlanilhaInvalida, match=trecho):
        _ler(dados, nome)


def test_numero_do_excel_nao_vira_decimal():
    """LEAD ID e telefone guardados como número no Excel não voltam com ".0"."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["LEAD ID", "NOME LEAD", "TELEFONE"])
    ws.append([12345678, "Lanchonete", 11988881234])
    buf = io.BytesIO()
    wb.save(buf)
    v = _ler(buf.getvalue(), "x.xlsx").validos[0]
    assert (v["codigo_externo"], v["phone"]) == ("12345678", "11988881234")


def test_sql_e_funil_tem_os_mesmos_status():
    """O CHECK do banco (001 e 005) e o LEGADO do 005 batem com app/funil.py.

    O SQL é escrito à mão; sem este teste, um status novo no Python entraria
    na tela e seria recusado pelo banco na primeira gravação.
    """
    sql = Path(__file__).resolve().parents[1] / "app" / "sql"
    for arquivo in ("001_schema.sql", "005_funil_campo.sql"):
        texto = (sql / arquivo).read_text(encoding="utf-8")
        # A regra é o CHECK; o 005 também tem um "stage IN (...)" com as etapas
        # antigas, no UPDATE da conversão.
        trecho = re.search(r"CHECK\s*\(stage\s+IN\s*\(([^)]*)\)", texto, re.S | re.I).group(1)
        assert tuple(re.findall(r"'(\w+)'", trecho)) == funil.CHAVES, arquivo
    legado = dict(re.findall(r"WHEN '(\w+)' THEN '(\w+)'",
                             (sql / "005_funil_campo.sql").read_text(encoding="utf-8")))
    assert legado == funil.LEGADO
